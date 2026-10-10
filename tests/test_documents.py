"""Document readers and the upload/watched-folder sync (needs Postgres for sync tests)."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import Chunk, Document
from app.documents.readers import EmptyDocument, UnsupportedDocument, read_document
from app.documents.service import DocumentError, safe_filename, save_upload, sync_documents, validate_watch_path

# A tiny valid one-page PDF containing the text "Parking costs 40 rupees per hour".
PDF_TEXT = "Parking costs 40 rupees per hour"


def make_pdf(text: str = PDF_TEXT) -> bytes:
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_read_pdf(tmp_path: Path) -> None:
    f = tmp_path / "parking.pdf"
    f.write_bytes(make_pdf())
    assert PDF_TEXT in read_document(f)


def test_read_docx_and_xlsx(tmp_path: Path) -> None:
    import docx
    from openpyxl import Workbook

    d = docx.Document()
    d.add_paragraph("Valet parking is available at Gate 2.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Floor"
    table.rows[0].cells[1].text = "Second"
    d.save(tmp_path / "info.docx")
    text = read_document(tmp_path / "info.docx")
    assert "Valet parking" in text and "Floor | Second" in text

    wb = Workbook()
    ws = wb.active
    ws.title = "Stores"
    ws.append(["Store", "Floor", "Phone"])
    ws.append(["ASICS", "Second Floor", "4054060"])
    wb.save(tmp_path / "stores.xlsx")
    assert "Store: ASICS | Floor: Second Floor | Phone: 4054060" in read_document(tmp_path / "stores.xlsx")


def test_read_text_encodings_and_errors(tmp_path: Path) -> None:
    (tmp_path / "a.md").write_text("# Hours\nലുലു 10 AM", encoding="utf-8")
    assert "ലുലു" in read_document(tmp_path / "a.md")
    (tmp_path / "b.txt").write_bytes("café".encode("cp1252"))
    assert read_document(tmp_path / "b.txt") == "café"
    (tmp_path / "empty.txt").write_text("   ")
    with pytest.raises(EmptyDocument):
        read_document(tmp_path / "empty.txt")
    (tmp_path / "bad.pdf").write_bytes(b"not a pdf")
    with pytest.raises(UnsupportedDocument):
        read_document(tmp_path / "bad.pdf")
    with pytest.raises(UnsupportedDocument):
        read_document(tmp_path / "x.exe")


def test_safe_filename() -> None:
    assert safe_filename("../../etc/passwd") == "passwd"
    assert safe_filename("C:\\docs\\Menu (1).pdf") == "Menu (1).pdf"
    assert safe_filename("മെനു.pdf") == "മെനു.pdf"


# ----------------------------------------------------------------------------- sync (DB)


@pytest.fixture
def data_dirs(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    from app.config import get_settings

    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("WATCHED_DIR", str(tmp_path / "watched"))
    (tmp_path / "watched").mkdir()
    get_settings.cache_clear()
    return tmp_path


@pytest.mark.integration
def test_upload_index_change_and_delete(db: Session, embedder, make_client, data_dirs: Path) -> None:
    client = make_client()
    doc = save_upload(db, client, "parking.pdf", io.BytesIO(make_pdf()))
    assert doc.status == "pending"

    stats = sync_documents(db, client, embedder)
    assert stats["indexed"] == 1
    db.refresh(doc)
    assert doc.status == "indexed" and doc.chunk_count >= 1
    assert db.scalar(select(Chunk).where(Chunk.document_id == doc.id)).source == "parking.pdf"

    # Unchanged file: skipped.
    assert sync_documents(db, client, embedder)["unchanged"] == 1
    # Changed content: re-indexed.
    save_upload(db, client, "parking.pdf", io.BytesIO(make_pdf("Parking is free on Sundays")))
    assert sync_documents(db, client, embedder)["indexed"] == 1
    assert "free on Sundays" in db.scalar(select(Chunk.content).where(Chunk.client_id == client.id))
    # File deleted from disk: row and chunks removed.
    (data_dirs / "data" / "clients" / client.client_id / "documents" / "parking.pdf").unlink()
    assert sync_documents(db, client, embedder)["removed"] == 1
    assert db.scalar(select(Chunk).where(Chunk.client_id == client.id)) is None


@pytest.mark.integration
def test_upload_rejects_bad_types(db: Session, make_client, data_dirs: Path) -> None:
    client = make_client()
    with pytest.raises(DocumentError):
        save_upload(db, client, "virus.exe", io.BytesIO(b"MZ"))
    with pytest.raises(DocumentError):
        save_upload(db, client, "empty.txt", io.BytesIO(b""))


@pytest.mark.integration
def test_watched_folder(db: Session, embedder, make_client, data_dirs: Path) -> None:
    folder = data_dirs / "watched" / "acme" / "sub"
    folder.mkdir(parents=True)
    (folder / "hours.txt").write_text("The cinema opens at 9 AM.")
    (folder / "ignore.exe").write_bytes(b"x")
    client = make_client()

    path = validate_watch_path("acme")  # relative to WATCHED_DIR
    client.document_settings = {**client.document_settings, "watch_path": path}
    db.commit()

    stats = sync_documents(db, client, embedder, sources=("folder",))
    assert stats["indexed"] == 1
    doc = db.scalar(select(Document).where(Document.client_id == client.id))
    assert doc.source == "folder" and doc.path == "sub/hours.txt"
    assert client.document_settings["last_scan_at"]

    with pytest.raises(DocumentError):
        validate_watch_path("/etc")  # outside WATCHED_DIR
    with pytest.raises(DocumentError):
        validate_watch_path("../")
    with pytest.raises(DocumentError):
        validate_watch_path("does-not-exist")


def make_scanned_pdf(path: Path, page_texts: list[str]) -> None:
    """An image-only PDF (no text layer), like a scanner produces."""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.load_default(size=48)
    pages = []
    for text in page_texts:
        img = Image.new("L", (2480, 3508), 255)  # A4 at 300 dpi
        ImageDraw.Draw(img).text((200, 300), text, fill=0, font=font)
        pages.append(img)
    pages[0].save(path, save_all=True, append_images=pages[1:], resolution=300)


@pytest.mark.skipif(not __import__("shutil").which("tesseract"), reason="tesseract not installed")
def test_scanned_pdf_is_ocrd_in_page_order(tmp_path: Path) -> None:
    f = tmp_path / "scan.pdf"
    make_scanned_pdf(f, [f"Parking costs {n} rupees per hour" for n in (10, 20, 30, 40, 50)])
    seen: list[tuple[int, int]] = []
    text = read_document(f, progress=lambda done, total: seen.append((done, total)))
    positions = [text.index(f"Parking costs {n} rupees") for n in (10, 20, 30, 40, 50)]
    assert positions == sorted(positions)
    assert "[Page 5]" in text
    assert seen[-1] == (5, 5)


def test_scanned_pdf_reading_can_be_cancelled(tmp_path: Path) -> None:
    from app.documents.readers import ReadCancelled

    f = tmp_path / "scan.pdf"
    make_scanned_pdf(f, ["Page one", "Page two"])
    with pytest.raises(ReadCancelled):
        read_document(f, should_stop=lambda: True)


@pytest.mark.integration
def test_document_deleted_while_indexing_is_skipped(db: Session, embedder, make_client, data_dirs: Path) -> None:
    from app.db.session import new_session
    from app.documents.service import delete_document

    client = make_client()
    gone = save_upload(db, client, "a-gone.pdf", io.BytesIO(make_pdf("Delete me while indexing")))
    kept = save_upload(db, client, "b-kept.pdf", io.BytesIO(make_pdf()))
    gone_id = gone.id

    class DeletingEmbedder:
        """Deletes the first document (as an admin would, in another session) while it is being embedded."""

        name, dim, count_tokens = embedder.name, embedder.dim, embedder.count_tokens

        def encode(self, texts):
            if any("Delete me" in t for t in texts):
                other = new_session()
                delete_document(other, other.get(Document, gone_id), other.get(type(client), client.id))
                other.close()
            return embedder.encode(texts)

    logs: list[str] = []
    stats = sync_documents(db, client, DeletingEmbedder(), log=logs.append)
    assert stats["indexed"] == 1 and any("deleted while being indexed" in line for line in logs)
    db.refresh(kept)
    assert kept.status == "indexed" and db.get(Document, gone_id) is None


@pytest.mark.integration
def test_disabled_document_is_not_used_for_answers(db: Session, embedder, make_client, data_dirs: Path) -> None:
    from app.search.hybrid import search

    client = make_client()
    doc = save_upload(db, client, "parking.pdf", io.BytesIO(make_pdf()))
    sync_documents(db, client, embedder)
    assert any(h.source == "parking.pdf" for h in search(db, embedder, client.id, "parking rupees per hour").hits)

    doc.enabled = False
    db.commit()
    assert not search(db, embedder, client.id, "parking rupees per hour").hits  # neither vector nor keyword hits
    db.refresh(doc)
    assert doc.status == "indexed" and doc.chunk_count >= 1  # still indexed: switching back is instant

    doc.enabled = True
    db.commit()
    assert search(db, embedder, client.id, "parking rupees per hour").hits
