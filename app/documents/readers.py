"""Extract plain text from supported document types (scanned PDF pages via Tesseract OCR)."""

from __future__ import annotations

import io
import logging
import os
import re
import shutil
import subprocess
import time
from collections import deque
from collections.abc import Callable
from concurrent.futures import Future, ThreadPoolExecutor
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".md"}
MAX_TEXT_CHARS = 50_000_000  # ~15,000 dense pages; only guards against absurd files

MIN_PAGE_CHARS = 50  # a PDF page with less text than this is treated as a scan and OCR'd
OCR_PAGE_TIMEOUT = 180  # seconds
PROGRESS_INTERVAL = 0.5  # seconds between progress reports (each one is a database write)

Progress = Callable[[int, int], None]  # (pages read, total pages)
_LEADER_RE = re.compile(r"(?:[ \t]*[.·•_…][ \t]*){5,}")  # dot leaders in contents pages


class UnsupportedDocument(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


class ReadCancelled(Exception):
    """The job was cancelled while a long document was being read."""


def is_supported(path: Path | str) -> bool:
    return Path(path).suffix.lower() in SUPPORTED_EXTENSIONS


def ocr_languages() -> str:
    from app.config import get_settings

    langs = get_settings().ocr_langs.strip() or "eng"
    if not re.fullmatch(r"[A-Za-z_]+(\+[A-Za-z_]+)*", langs):
        raise UnsupportedDocument(f"OCR_LANGS is invalid: {langs!r} (expected e.g. eng or eng+mal).")
    return langs


def ocr_dpi() -> int:
    from app.config import get_settings

    return get_settings().ocr_dpi


def ocr_workers() -> int:
    from app.config import get_settings

    return get_settings().ocr_workers or os.cpu_count() or 2


def ocr_image(image: Any, langs: str) -> str:
    """OCR one page image with the tesseract binary (runs outside the GIL, so threads parallelise)."""
    buf = io.BytesIO()
    image.save(buf, format="PPM")  # uncompressed: far quicker to hand over than PNG
    result = subprocess.run(
        ["tesseract", "-", "-", "-l", langs, "--psm", "3"],
        input=buf.getvalue(),
        capture_output=True,
        timeout=OCR_PAGE_TIMEOUT,
        # One core per page: parallelism comes from OCR'ing several pages at once.
        env={**os.environ, "OMP_THREAD_LIMIT": "1"},
        check=False,
    )
    if result.returncode != 0:
        message = result.stderr.decode("utf-8", "replace").strip().splitlines()
        raise UnsupportedDocument(f"OCR failed: {message[-1] if message else 'tesseract error'}")
    return result.stdout.decode("utf-8", "replace").strip()


def read_pdf(path: Path, progress: Progress | None = None, should_stop: Callable[[], bool] | None = None) -> str:
    """Text per page with PDFium; pages without a text layer (scans) are OCR'd in parallel."""
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(str(path))
    except pdfium.PdfiumError as exc:
        if "password" in str(exc).lower():
            raise UnsupportedDocument("The PDF is password-protected.") from exc
        raise
    total = len(pdf)
    texts: list[str] = [""] * total
    can_ocr = shutil.which("tesseract") is not None
    langs = ocr_languages() if can_ocr else ""
    workers = ocr_workers()
    scale = ocr_dpi() / 72
    pool: ThreadPoolExecutor | None = None
    in_flight: deque[tuple[int, Future[str]]] = deque()
    done = 0
    last_report = 0.0

    def report(force: bool = False) -> None:
        nonlocal last_report
        now = time.monotonic()
        if progress and (force or now - last_report >= PROGRESS_INTERVAL):
            last_report = now
            progress(done, total)

    def collect_oldest() -> None:
        nonlocal done
        index, future = in_flight.popleft()
        ocr_text = future.result()
        if len(ocr_text) > len(texts[index]):
            texts[index] = ocr_text
        done += 1
        report()

    try:
        for index in range(total):
            if should_stop and should_stop():
                raise ReadCancelled()
            page = pdf[index]
            try:
                textpage = page.get_textpage()
                try:
                    texts[index] = textpage.get_text_bounded().strip()
                finally:
                    textpage.close()
                if len(texts[index]) >= MIN_PAGE_CHARS or not can_ocr:
                    done += 1
                    report()
                    continue
                # PDFium is not thread-safe, so pages are rendered here and only Tesseract runs in the pool.
                image = page.render(scale=scale, grayscale=True).to_pil()
            finally:
                page.close()
            if pool is None:
                pool = ThreadPoolExecutor(max_workers=workers, thread_name_prefix="ocr")
            in_flight.append((index, pool.submit(ocr_image, image, langs)))
            # Bound the rendered pages held in memory (~4 MB each at 200 dpi, ~9 MB at 300).
            while len(in_flight) >= workers * 2:
                collect_oldest()
        while in_flight:
            if should_stop and should_stop():
                raise ReadCancelled()
            collect_oldest()
        report(force=True)
    finally:
        if pool is not None:
            pool.shutdown(wait=True, cancel_futures=True)
        pdf.close()
    if not can_ocr and not any(texts):
        logger.warning("%s looks scanned but tesseract is not installed; OCR skipped", path.name)
    return "\n\n".join(f"[Page {i + 1}]\n{t}" for i, t in enumerate(texts) if t)


def read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            # Merged cells repeat their text; keep each value once per row.
            unique = list(dict.fromkeys(c for c in cells if c))
            if unique:
                parts.append(" | ".join(unique))
    return "\n".join(parts)


def read_xlsx(path: Path) -> str:
    """Each row becomes 'Header: value | Header: value' so facts stay self-describing."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header: list[str] | None = None
            lines: list[str] = []
            for row in rows:
                values = ["" if v is None else str(v).strip() for v in row]
                if not any(values):
                    continue
                if header is None:
                    header = values
                    continue
                pairs = [f"{header[i] or f'Column {i + 1}'}: {v}" if i < len(header) else v for i, v in enumerate(values) if v]
                lines.append(" | ".join(pairs))
            if header is not None and not lines:  # single-row sheet
                lines.append(" | ".join(v for v in header if v))
            if lines:
                parts.append(f"[Sheet: {sheet.title}]\n" + "\n".join(lines))
    finally:
        workbook.close()
    return "\n\n".join(parts)


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):  # UTF-16 only when it has a byte-order mark
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


READERS = {".pdf": read_pdf, ".docx": read_docx, ".xlsx": read_xlsx, ".txt": read_text, ".md": read_text}


def read_document(path: Path, progress: Progress | None = None, should_stop: Callable[[], bool] | None = None) -> str:
    """Return the document's text. Raises UnsupportedDocument / EmptyDocument with a user-facing message.

    ``progress(pages_read, total_pages)`` and ``should_stop`` apply to PDFs, which can take minutes to OCR.
    """
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise UnsupportedDocument(f"Unsupported file type: {path.suffix or 'none'}")
    try:
        text = read_pdf(path, progress, should_stop) if reader is read_pdf else reader(path)
    except (UnsupportedDocument, EmptyDocument, ReadCancelled):
        raise
    except Exception as exc:  # noqa: BLE001 - corrupt files raise all kinds of errors
        raise UnsupportedDocument(f"Could not read the file ({exc.__class__.__name__}). Is it corrupt?") from exc
    # Table-of-contents leaders ("Overview ........ 3") drown out the real words of a chunk in search.
    text = _LEADER_RE.sub(" … ", text).strip()
    if not text:
        hint = ""
        if path.suffix.lower() == ".pdf":
            hint = (
                " OCR found no readable text either: the scan may be blank or too low-resolution."
                if shutil.which("tesseract")
                else " It is a scanned PDF and OCR (tesseract) is not installed on this server."
            )
        raise EmptyDocument("No text found in the document." + hint)
    return text[:MAX_TEXT_CHARS]
